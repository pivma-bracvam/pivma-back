"""Arquivo-modelo do template de coleta (Spec 041, FR-012 a FR-016)."""

import io
from types import SimpleNamespace

from openpyxl import load_workbook

from pivma.core.collection_template_service import (
    file_header,
    render_csv,
    render_xlsx,
)

FIXED = ['codigo_amostra', 'experimento', 'replica']


def _column(key, position):
    return SimpleNamespace(key=key, position=position)


def test_header_has_fixed_columns_then_keys_by_position():
    columns = [_column('lote', 7), _column('viabilidade', 2), _column('d', 5)]

    assert file_header(columns) == [*FIXED, 'viabilidade', 'd', 'lote']


def test_header_without_columns_has_only_fixed_columns():
    assert file_header([]) == FIXED


def test_csv_has_bom_semicolons_and_a_single_line():
    content = render_csv([*FIXED, 'viabilidade'])

    assert content.startswith(b'\xef\xbb\xbf')
    lines = content.decode('utf-8-sig').splitlines()
    assert lines == ['codigo_amostra;experimento;replica;viabilidade']


def test_xlsx_has_one_sheet_with_header_only_and_no_validation():
    header = [*FIXED, 'viabilidade']

    workbook = load_workbook(io.BytesIO(render_xlsx(header)))

    assert workbook.sheetnames == ['resultados']
    sheet = workbook['resultados']
    assert [cell.value for cell in sheet[1]] == header
    assert sheet.max_row == 1
    assert sheet.data_validations.dataValidation == []
